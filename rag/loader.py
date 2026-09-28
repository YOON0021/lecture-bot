"""강의 자료(PDF, PPTX)를 페이지/슬라이드 단위로 읽어서 청크로 나눈다."""

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

SUPPORTED = {".pdf", ".pptx"}


@dataclass
class Chunk:
    id: str
    text: str
    subject: str  # 과목
    source: str  # 파일명
    page: int  # PDF는 페이지, PPTX는 슬라이드 번호 (1부터)


def _clean(text: str) -> str:
    return " ".join(text.split())


def load_pdf(path: Path) -> list[tuple[int, str]]:
    pages = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            pages.append((i, text))
    return pages


def load_pptx(path: Path) -> list[tuple[int, str]]:
    slides = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide:  # 발표자 노트에 설명이 들어있는 경우가 많다
            parts.append(slide.notes_slide.notes_text_frame.text)
        text = _clean(" ".join(parts))
        if text:
            slides.append((i, text))
    return slides


def split_text(text: str, chunk_size: int = 800, overlap: int = 150) -> list[str]:
    """글자 수 기준으로 자르되, 가능하면 문장 끝(. ? !)에서 자른다."""
    if len(text) <= chunk_size:
        return [text]

    chunks, start = [], 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            cut = max(text.rfind(p, start + chunk_size // 2, end) for p in (". ", "? ", "! "))
            if cut != -1:
                end = cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def load_pages(path: str | Path) -> list[tuple[int, str]]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".pptx":
        return load_pptx(path)
    raise ValueError(f"지원하지 않는 형식입니다: {ext}")


def chunk_file(path: str | Path, subject: str, name: str | None = None) -> list[Chunk]:
    """청크가 페이지 경계를 넘지 않게 해서 출처 페이지를 정확히 표시할 수 있게 한다."""
    name = name or Path(path).name
    chunks = []
    for page_no, text in load_pages(path):
        for j, piece in enumerate(split_text(text)):
            chunks.append(
                Chunk(id=f"{subject}/{name}:p{page_no}:c{j}", text=piece, subject=subject, source=name, page=page_no)
            )
    return chunks
